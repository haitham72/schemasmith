"""Synthetic WP6 checks; real Postgres parity remains opt-in and requires a real fixture."""
import copy
from datetime import date, timedelta
import importlib.util
import io
import json
import os
from pathlib import Path
import sys

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = ROOT / 'app' / 'export.py'
if MODULE_PATH.exists():
    loader = importlib.util.spec_from_file_location('exports_under_test', MODULE_PATH)
    exports = importlib.util.module_from_spec(loader)
    sys.modules[loader.name] = exports
    loader.loader.exec_module(exports)
else:
    from app import export as exports


class Cursor:
    def __init__(self, rows):
        self.rows = rows
    def fetchall(self):
        return copy.deepcopy(self.rows)


class MemoryConnection:
    def __init__(self, tables):
        self.tables, self.statements = tables, []
    def execute(self, statement):
        rendered = statement.as_string()
        self.statements.append(rendered)
        assert rendered.startswith('select * from "')
        table = rendered.split('"')[1]
        return Cursor(self.tables[table])


@pytest.fixture
def example_a():
    proposal = json.loads((Path(__file__).resolve().parent / 'fixtures' / 'exports_configuration.json').read_text(encoding='utf-8'))
    spec = {'entities':[
        {'name':'topic','kind':'reference','hierarchy':{'max_depth':4},'fields':[{'name':'name','type':'text'}]},
        {'name':'tool','kind':'reference','parent':'topic','fields':[{'name':'name','type':'text'},{'name':'domains','type':'list<domain>'},{'name':'needs_review','type':'bool'}]},
        {'name':'request','kind':'record','owner':'requester','fields':[{'name':'reason','type':'text'}],'lines':{'entity_ref':'tool'}}],
        'workflow':{'entity':'request','level':'line','initial':'pending','transitions':[{'collects':'justification'}],
            'rollup':{'all':{'approved':'approved','denied':'denied'},'mixed':'partial','none_decided':'pending'}},
        'views':[dict(type='report',entity='request',formats=['csv','xlsx'],**proposal['report_view_additions'])],
        'exports':[dict(name='allowlist',**{'from':'request.lines.tool.domains'},distinct=True,where={'state':'approved'},formats=['txt','csv','xlsx'],**proposal['export_additions'])],
        'dashboard':[{'label':'Waiting','count':'request','where':{'state':'pending'}},{'label':'Mixed','count':'request','where':{'state':'partial'}}]}
    tables={'users':[{'id':'u','name':'Test User','email':'user@example.test','status':'approved'}],
        'x_topic':[{'id':'t','name':'Development'},{'id':'sub','name':'Packages','parent_id':'t'}],
        'x_tool':[{'id':'a','parent_id':'sub','name':'PyPI (all packages)','domains':['shared.example','pypi.example'],'needs_review':False},
            {'id':'b','parent_id':'sub','name':'npm (all packages)','domains':['shared.example'],'needs_review':True},
            {'id':'c','parent_id':'sub','name':'Untouched','domains':['untouched.example'],'needs_review':False}],
        'x_request':[{'id':'r','owner_id':'u','reason':'Build packages','created_at':'2026-01-02','active':True},
                     {'id':'old','owner_id':'u','reason':'Earlier','created_at':'2026-01-01','active':True}],
        'x_request_lines':[{'id':'la','record_id':'r','ref_id':'a','state':'approved','justification':''},
            {'id':'lb','record_id':'r','ref_id':'b','state':'denied','justification':'Consider alternatives'},
            {'id':'lc','record_id':'old','ref_id':'b','state':'approved','justification':''}]}
    return spec, tables


@pytest.fixture
def example_b():
    spec={'entities':[{'name':'asset','kind':'reference','fields':[{'name':'name','type':'text'}]},
        {'name':'checkout','kind':'record','fields':[{'name':'asset','type':'ref','to':'asset'},{'name':'due_date','type':'date'},{'name':'purpose','type':'text'}]}],
        'workflow':{'entity':'checkout','level':'record','initial':'requested','transitions':[]},
        'views':[{'name':'loans','type':'report','entity':'checkout','columns':[{'source':'asset.name','label':'Equipment'},{'source':'due_date','label':'Due'},{'source':'purpose','label':'Purpose'}],'formats':['csv','xlsx']}],
        'exports':[{'name':'overdue','from':'checkout','where':{'state':'checked_out','due_date':'<today'},'formats':['csv','xlsx']}],
        'dashboard':[{'label':'Out','count':'checkout','where':{'state':'checked_out'}}]}
    tables={'users':[], 'x_asset':[{'id':'a','name':'Camera'}],
        'x_checkout':[{'id':'c','asset':'a','due_date':date.today()-timedelta(days=1),'purpose':'=HYPERLINK("bad")','state':'checked_out'}]}
    return spec,tables


def test_report_a_workbook_preserves_original_assertions(example_a):
    spec,tables=example_a
    result=exports.export_report(MemoryConnection(tables),spec,'approvals','xlsx')
    wb=openpyxl.load_workbook(io.BytesIO(result.content))
    assert wb.sheetnames == ['Test User','Reviewer_Matrix','Summary']
    ws=wb['Test User']
    assert [c.value for c in ws[2]] == ['Category Path','Item Name','Target Domains','Business Reason','Status','Reviewer Notes']
    assert ws['A1'].value.startswith('ACCESS REQUEST CATALOG')
    values={ws.cell(r,2).value:ws.cell(r,5).value for r in range(3,ws.max_row+1)}
    assert values == {'PyPI (all packages)':'Approved','npm (all packages)':'Denied'}
    assert len(ws.data_validations.dataValidation)==1
    assert 'Justify' in ws.data_validations.dataValidation[0].formula1
    assert ws['A3'].value=='Development > Packages'
    assert ws.freeze_panes=='A3'
    assert ws['A1'].fill.fgColor.rgb.endswith('2C3E50')
    assert len(ws.conditional_formatting)==1
    matrix=wb['Reviewer_Matrix']
    assert [c.value for c in matrix[2]][:3]==['Category Path','Item Name','Domain']
    assert matrix['E3'].data_type=='f'
    assert matrix['E3'].value.startswith('=COUNTIF(')
    assert {matrix.cell(r,4).value for r in (3,4)}=={'ALLOW','DENY'}
    summary=wb['Summary']
    metrics={summary.cell(r,1).value:summary.cell(r,2).value for r in range(3,11)}
    assert metrics['Tools requested']==2 and metrics['Approved']==1 and metrics['Denied']==1


def test_provenance_and_boolean_group_aggregate(example_a):
    spec,tables=example_a
    result=exports.export_declared(MemoryConnection(tables),spec,'allowlist','csv')
    assert result.content == b'domain,granted_by,verified\r\npypi.example,PyPI (all packages),yes\r\nshared.example,PyPI (all packages); npm (all packages),no\r\n'
    assert exports.export_declared(MemoryConnection(tables),spec,'allowlist','txt').content == b'pypi.example\nshared.example\n'


@pytest.mark.parametrize('fixture,name',[('example_a','approvals'),('example_b','loans')])
@pytest.mark.parametrize('fmt',['csv','xlsx'])
def test_both_report_formats(request,fixture,name,fmt):
    spec,tables=request.getfixturevalue(fixture)
    output=exports.export_report(MemoryConnection(tables),spec,name,fmt)
    assert output.content and output.filename.endswith('.'+fmt)
    if fmt=='xlsx':
        wb=openpyxl.load_workbook(io.BytesIO(output.content))
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    if cell.data_type=='f':
                        assert cell.value.startswith('=COUNTIF(')


@pytest.mark.parametrize('fmt',['csv','xlsx'])
def test_second_domain_declared_export(example_b,fmt):
    spec,tables=example_b
    output=exports.export_declared(MemoryConnection(tables),spec,'overdue',fmt)
    assert output.content and output.filename=='overdue.'+fmt
    if fmt=='csv':
        assert b"'=HYPERLINK" in output.content


def test_record_counts_not_line_counts(example_a,example_b):
    spec,tables=example_a
    assert exports.dashboard(MemoryConnection(tables),spec)==[{'label':'Waiting','value':0},{'label':'Mixed','value':1}]
    spec,tables=example_b
    assert exports.dashboard(MemoryConnection(tables),spec)==[{'label':'Out','value':1}]


@pytest.mark.parametrize('value',['=1+1','+cmd','-cmd','@SUM(A1)','  =1+1','\t=cmd'])
def test_user_values_are_never_excel_formulas(example_b,value):
    spec,tables=example_b
    tables['x_checkout'][0]['purpose']=value
    output=exports.export_report(MemoryConnection(tables),spec,'loans','xlsx')
    cell=openpyxl.load_workbook(io.BytesIO(output.content)).active['C3']
    assert cell.data_type=='s' and cell.value=="'"+value


def test_unknown_selectors_fail_closed(example_b):
    spec,tables=example_b
    spec['views'][0]['columns']=[{'source':'purpose; drop table users','label':'Bad'}]
    with pytest.raises(exports.ExportError,match='Unknown selector'):
        exports.export_report(MemoryConnection(tables),spec,'loans','csv')


def test_distinct_requires_explicit_aggregation(example_a):
    spec,tables=example_a
    spec['exports'][0]['columns'][1]['aggregate']='value'
    with pytest.raises(exports.ExportError,match='explicit aggregate'):
        exports.export_declared(MemoryConnection(tables),spec,'allowlist','csv')


def test_archived_roots_excluded_historical_refs_retained(example_a):
    spec,tables=example_a
    tables['x_request'][0]['active']=False
    tables['x_tool'][1]['active']=False
    assert exports.export_declared(MemoryConnection(tables),spec,'allowlist','txt').content == b'shared.example\n'


def test_zero_rows_workbook_is_valid(example_a):
    spec,tables=example_a
    tables['x_request']=[]
    output=exports.export_report(MemoryConnection(tables),spec,'approvals','xlsx')
    wb=openpyxl.load_workbook(io.BytesIO(output.content))
    assert wb.sheetnames==['Reviewer_Matrix','Summary']


def test_sheet_names_casefold_and_reserved_names(example_a):
    spec,tables=example_a
    tables['users'][0]['name']='Summary'
    output=exports.export_report(MemoryConnection(tables),spec,'approvals','xlsx')
    assert openpyxl.load_workbook(io.BytesIO(output.content)).sheetnames==['Summary (2)','Reviewer_Matrix','Summary']


def test_queries_quote_identifiers(example_b):
    spec,tables=example_b
    conn=MemoryConnection(tables)
    exports.export_report(conn,spec,'loans','csv')
    assert conn.statements==['select * from "x_asset"','select * from "x_checkout"','select * from "users"']


@pytest.mark.skipif(not os.getenv('EXPORTS_PARITY_DATABASE_URL'),reason='Owner deferred real PostgreSQL; requires seeded baseline DB')
def test_postgres_baseline():
    import psycopg
    from psycopg.rows import dict_row
    from app.spec import validate
    from app.migrate import active
    candidates=[Path(__file__).resolve().parent/'fixtures'/'allowlist_baseline.txt']
    fixture=next((item for item in candidates if item.exists()),None)
    if fixture is None:
        pytest.skip('baseline-not-captured')
    with psycopg.connect(os.environ['EXPORTS_PARITY_DATABASE_URL'],row_factory=dict_row) as conn:
        spec=validate(active(conn)['spec'])
        result=exports.export_declared(conn,spec,'allowlist','csv')
        assert result.content==fixture.read_bytes()


def test_allowlist_baseline_fixture_state():
    """No baseline bytes are fabricated here: when
    tests/fixtures/allowlist_baseline.txt is absent, record the gap and skip."""
    candidates=[Path(__file__).resolve().parent/'fixtures'/'allowlist_baseline.txt']
    found=[str(item) for item in candidates if item.exists()]
    if not found:
        pytest.skip('baseline-not-captured')
    assert found and all(Path(item).read_bytes() for item in found)


@pytest.mark.skipif(not os.getenv('EXPORTS_TEST_DATABASE_URL'), reason='Owner deferred real PostgreSQL')
@pytest.mark.parametrize('fixture,report,declared',[('example_a','approvals','allowlist'),('example_b','loans','overdue')])
def test_postgres_synthetic_both_examples(request,fixture,report,declared):
    """Temporary tables shadow public tables; the transaction always rolls back."""
    import psycopg
    from psycopg import sql
    from psycopg.rows import dict_row
    spec,tables=request.getfixturevalue(fixture)
    with psycopg.connect(os.environ['EXPORTS_TEST_DATABASE_URL'],row_factory=dict_row) as conn:
        with conn.transaction(force_rollback=True):
            for table,rows in tables.items():
                fields={key for row in rows for key in row} or {'id','name','email','status'}
                fields=sorted(fields)
                definitions=[]
                for field in fields:
                    sample=next((row[field] for row in rows if row.get(field) is not None), '')
                    kind='boolean' if isinstance(sample,bool) else 'date' if type(sample) is date else 'text[]' if isinstance(sample,list) else 'text'
                    definitions.append(sql.Composed([sql.Identifier(field),sql.SQL(' '+kind)]))
                conn.execute(sql.Composed([sql.SQL('create temporary table '),sql.Identifier(table),sql.SQL(' ('),sql.SQL(', ').join(definitions),sql.SQL(') on commit drop')]))
                for row in rows:
                    conn.execute(sql.Composed([sql.SQL('insert into '),sql.Identifier(table),sql.SQL(' ('),sql.SQL(', ').join(map(sql.Identifier,fields)),sql.SQL(') values ('),sql.SQL(', ').join([sql.Placeholder()]*len(fields)),sql.SQL(')')]),[row.get(field) for field in fields])
            for fmt in ('csv','xlsx'):
                assert exports.export_report(conn,spec,report,fmt).content
                assert exports.export_declared(conn,spec,declared,fmt).content
            assert exports.dashboard(conn,spec)


def test_provenance_distinct_by_preserves_duplicate_names(example_a):
    spec,tables=example_a
    tables['x_tool'][1]['name']=tables['x_tool'][0]['name']
    output=exports.export_declared(MemoryConnection(tables),spec,'allowlist','csv')
    assert b'shared.example,PyPI (all packages); PyPI (all packages),no' in output.content


def test_per_person_summary_section_retained(example_a):
    spec,tables=example_a
    output=exports.export_report(MemoryConnection(tables),spec,'approvals','xlsx')
    sheet=openpyxl.load_workbook(io.BytesIO(output.content))['Summary']
    assert sheet['A12'].value=='PER PERSON'
    assert [cell.value for cell in sheet[13]]==['Person','Tickets','Tools','Categories','Approved','Denied','Justify','Pending']
    assert [cell.value for cell in sheet[14]]==['Test User',2,2,1,1,1,0,0]


def spec_module():
    name='spec_under_test'
    if name in sys.modules:
        return sys.modules[name]
    path=ROOT/'app'/'spec.py'
    loader=importlib.util.spec_from_file_location(name,path)
    module=importlib.util.module_from_spec(loader)
    sys.modules[name]=module
    loader.loader.exec_module(module)
    return module


@pytest.mark.parametrize('name',['access_requests','equipment_checkout'])
def test_examples_validate_and_round_trip(name):
    models=spec_module()
    document=json.loads((ROOT/'examples'/name/'spec.json').read_text(encoding='utf-8'))
    spec=models.validate(document)
    assert models.validate(spec.model_dump(mode='json',by_alias=True,exclude_unset=True))==spec


def test_schema_matches_models():
    models=spec_module()
    disk=json.loads((ROOT/'spec'/'schema.json').read_text(encoding='utf-8'))
    assert disk==models.schema()


def test_closed_models_reject_free_selectors_and_formulas():
    import copy
    models=spec_module()
    document=json.loads((ROOT/'examples'/'access_requests'/'spec.json').read_text(encoding='utf-8'))
    bad=copy.deepcopy(document)
    bad['views'][3]['columns'][0]['source']='request.lines.tool.unknown'
    with pytest.raises(models.SpecValidationError) as raised:
        models.validate(bad)
    assert any(error['code']=='unresolved_selector' for error in raised.value.errors)
    bad=copy.deepcopy(document)
    bad['exports'][0]['columns'][1]['aggregate']='sum'
    with pytest.raises(models.SpecValidationError):
        models.validate(bad)
    bad=copy.deepcopy(document)
    bad['views'][3]['workbook']['matrix']['count_value']='PEND'
    assert 'PEND' in bad['views'][3]['workbook']['matrix']['state_labels'].values()
    bad['views'][3]['workbook']['matrix']['count_value']='missing-label'
    with pytest.raises(models.SpecValidationError) as raised:
        models.validate(bad)
    assert any(error['code']=='invalid_matrix_count' for error in raised.value.errors)
    bad=copy.deepcopy(document)
    bad['views'][3]['workbook']['state_labels']['approved']='Has, comma'
    with pytest.raises(models.SpecValidationError) as raised:
        models.validate(bad)
    assert any(error['code']=='invalid_excel_label' for error in raised.value.errors)


def test_workbook_config_covers_review_gates():
    models=spec_module()
    document=json.loads((ROOT/'examples'/'access_requests'/'spec.json').read_text(encoding='utf-8'))
    view=next(view for view in document['views'] if view.get('name')=='approvals')
    assert view['workbook']['matrix']['count_value'] in view['workbook']['matrix']['state_labels'].values()
    assert {metric['kind'] for metric in view['workbook']['summary']['metrics']} >= {'owners','groups','items','records','state'}
    export=next(export for export in document['exports'] if export['name']=='allowlist')
    assert {column.get('aggregate', 'value') for column in export['columns']} <= {'value','join_distinct','none_true'}
    assert export['label']=='Firewall allowlist'
    loans=json.loads((ROOT/'examples'/'equipment_checkout'/'spec.json').read_text(encoding='utf-8'))
    report=next(view for view in loans['views'] if view.get('name')=='loans')
    assert set(report['formats'])=={'csv','xlsx'}
