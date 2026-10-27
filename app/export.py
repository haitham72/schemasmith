"""Spec-driven reports and exports, retaining the source workbook styling helpers."""
import io
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FILL_BANNER = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
FILL_HEADER = PatternFill(start_color="D9E2EC", end_color="D9E2EC", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
FILL_ALLOW = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
FILL_DENY = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")
FILL_PEND = PatternFill(start_color="FEF7E0", end_color="FEF7E0", fill_type="solid")

FONT_BANNER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="102A43")
FONT_BODY = Font(name="Calibri", size=10)

THIN_BORDER = Border(
    left=Side(style="thin", color="BCCCDC"),
    right=Side(style="thin", color="BCCCDC"),
    top=Side(style="thin", color="BCCCDC"),
    bottom=Side(style="thin", color="BCCCDC"),
)

def sheet_name_for(name: str, used: set) -> str:
    base = re.sub(r'[\\/*?:\[\]]', "", (name or "User").strip())[:31] or "User"
    title, n = base, 2
    while title.casefold() in {item.casefold() for item in used}:
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(title)
    return title


def _banner(ws, title: str, ncols: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=_safe(title))
    c.fill = FILL_BANNER
    c.font = FONT_BANNER
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 22


def _header(ws, headers: list, center=()) -> None:
    ws.append([_safe(value) for value in headers])
    for n, _h in enumerate(headers, 1):
        c = ws.cell(row=2, column=n)
        c.fill = FILL_HEADER
        c.font = FONT_HEADER
        c.alignment = Alignment(horizontal="center" if n in center else "left", vertical="center")
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}2"


def _widths(ws) -> None:
    for col in ws.columns:
        widest = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min(widest + 3, 60), 12)


def _body_row(ws, row_num: int, values: list, zebra: bool, center=()) -> None:
    for n, v in enumerate(values, 1):
        c = ws.cell(row=row_num, column=n, value=_safe(v))
        c.font = FONT_BODY
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal="center" if n in center else "left", vertical="center", wrap_text=True)
        if zebra:
            c.fill = FILL_ZEBRA


def _dropdown(ws, values: list, target: str) -> None:
    dv = DataValidation(type="list", formula1='"' + ','.join(values) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(target)



import csv
from dataclasses import dataclass
from datetime import date, datetime
import json
from psycopg import sql


class ExportError(ValueError):
    pass


class _Formula(str):
    """Only engine-created formulas use this private marker."""


def _safe(value):
    if isinstance(value, _Formula):
        return str(value)
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def _doc(value):
    return value.model_dump(mode='json', by_alias=True, exclude_none=True) if hasattr(value, 'model_dump') else value


def _text(value):
    if value is None:
        return ''
    if isinstance(value, (list, tuple)):
        return ', '.join(_text(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, default=str, ensure_ascii=False, sort_keys=True)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).lower() if isinstance(value, bool) else str(value)


@dataclass(frozen=True)
class Download:
    content: bytes
    mime: str
    filename: str


class Dataset:
    """Load one consistent transaction's data; no pagination can silently lose rows.

    Only validated entity names become quoted identifiers. Selectors are interpreted
    against this closed graph, never inserted into SQL. Archived references remain
    available for historical records; archived root records are excluded.
    """
    def __init__(self, conn, spec):
        self.spec = _doc(spec)
        self.entities = {item['name']: item for item in self.spec['entities']}
        self.tables = {}
        for name, entity in self.entities.items():
            self.tables[name] = self._read(conn, 'x_' + name)
            if entity.get('lines'):
                self.tables[name + '_lines'] = self._read(conn, 'x_' + name + '_lines')
        self.users = self._read(conn, 'users') if any(e['kind'] == 'record' for e in self.entities.values()) else []
        self.index = {name: {str(row['id']): row for row in rows} for name, rows in self.tables.items()}
        self.index['users'] = {str(row['id']): row for row in self.users}
        wf = self.spec.get('workflow')
        if wf and wf['level'] == 'line':
            rule = wf.get('rollup')
            for row in self.tables[wf['entity']]:
                states = [line['state'] for line in self.tables[wf['entity'] + '_lines']
                          if str(line['record_id']) == str(row['id'])]
                if rule:
                    row['state'] = (rule['none_decided'] if not any(s in rule['all'] for s in states)
                                    else rule['all'][states[0]] if len(set(states)) == 1 and states[0] in rule['all']
                                    else rule['mixed'])
                else:
                    row['state'] = states[0] if states and len(set(states)) == 1 else wf['initial']

    @staticmethod
    def _read(conn, table):
        return [dict(row) for row in conn.execute(sql.Composed([sql.SQL('select * from '), sql.Identifier(table)])).fetchall()]

    def entity(self, name):
        if name not in self.entities:
            raise ExportError('Unknown entity: ' + name)
        return self.entities[name]

    def rows(self, name):
        self.entity(name)
        return [row for row in self.tables[name] if row.get('active', True)]

    def contexts(self, name):
        entity = self.entity(name)
        for row in self.rows(name):
            lines = [line for line in self.tables.get(name + '_lines', []) if str(line['record_id']) == str(row['id'])]
            if entity.get('lines'):
                for line in lines:
                    yield {'root': row, 'entity': name, 'line': line}
            else:
                yield {'root': row, 'entity': name, 'line': None}

    def value(self, context, source):
        if source == '$value':
            return context.get('$value')
        parts = source.split('.')
        entity, row = self.entity(context['entity']), context['root']
        if parts and parts[0] == entity['name']:
            parts = parts[1:]
        line = False
        if parts and parts[0] == 'lines':
            if not entity.get('lines'):
                raise ExportError('Entity has no lines')
            row, line, parts = context['line'] or {}, True, parts[1:]
            target = entity['lines']['entity_ref']
            if parts and parts[0] in ('ref', target):
                row = self.index[target].get(str(row.get('ref_id')), {})
                entity, line, parts = self.entity(target), False, parts[1:]
        while parts:
            part, parts = parts[0], parts[1:]
            if part == 'owner' and entity['kind'] == 'record' and not line:
                if len(parts) != 1 or parts[0] not in ('name', 'email', 'id'):
                    raise ExportError('Owner selectors expose only id, name and email')
                return self.index['users'].get(str(row.get('owner_id')), {}).get(parts[0])
            if part in ('parent', 'parent_path') and not line:
                target = entity.get('parent') or (entity['name'] if entity.get('hierarchy') else None)
                if target is None:
                    raise ExportError('Entity has no parent')
                if part == 'parent_path':
                    if len(parts) != 1:
                        raise ExportError('parent_path must end with a label field')
                    label = parts[0]
                    if label not in {f['name'] for f in self.entity(target)['fields']}:
                        raise ExportError('Unknown path label')
                    path, seen = [], set()
                    parent = self.index[target].get(str(row.get('parent_id')))
                    while parent:
                        key = str(parent['id'])
                        if key in seen:
                            raise ExportError('Parent hierarchy contains a cycle')
                        seen.add(key)
                        path.append(_text(parent.get(label)))
                        parent = self.index[target].get(str(parent.get('parent_id'))) if self.entity(target).get('hierarchy') else None
                    return ' > '.join(reversed(path))
                row, entity = self.index[target].get(str(row.get('parent_id')), {}), self.entity(target)
                continue
            fields = {field['name']: field for field in entity['fields']} if not line else {}
            collected = {t.get(k) for t in self.spec.get('workflow', {}).get('transitions', []) for k in ('collects', 'requires')} if line else set()
            if part not in fields and part not in {'id', 'created_at', 'updated_at', 'state', 'sort', 'active', *collected}:
                raise ExportError('Unknown selector field: ' + part)
            value = row.get(part)
            if not parts:
                return value
            field = fields.get(part, {})
            if field.get('type') != 'ref':
                raise ExportError('Only scalar references may be traversed')
            entity = self.entity(field['to'])
            row = self.index[entity['name']].get(str(value), {})
        return row

    def matches(self, context, filters, *, record_state=False, today=None):
        today = today or date.today()
        for source, expected in filters.items():
            if source == 'state' and context['line'] is not None and not record_state:
                value = context['line'].get('state')
            else:
                value = self.value(context, source)
            if isinstance(expected, str) and expected in ('<today', '<=today', '>today', '>=today'):
                if value is None:
                    return False
                if isinstance(value, datetime):
                    value = value.date()
                elif isinstance(value, str):
                    try:
                        value = date.fromisoformat(value)
                    except ValueError as exc:
                        raise ExportError('Relative date filter requires a date') from exc
                if type(value) is not date:
                    raise ExportError('Relative date filter requires a date')
                if not {'<today': value < today, '<=today': value <= today,
                        '>today': value > today, '>=today': value >= today}[expected]:
                    return False
            elif value != expected:
                return False
        return True


def _view(spec, view):
    spec = _doc(spec)
    if isinstance(view, str):
        matches = [v for v in spec.get('views', []) if v.get('name') == view and v['type'] == 'report']
        if len(matches) != 1:
            raise ExportError('Unknown report')
        return matches[0]
    return _doc(view)


def _columns(ds, view):
    columns = view.get('columns')
    if columns is None:
        columns = [field['name'] for field in ds.entity(view['entity'])['fields']]
    result = []
    fields = {field['name']: field for field in ds.entity(view['entity'])['fields']}
    for column in columns:
        result.append({'source': column, 'label': fields.get(column, {}).get('label', column)} if isinstance(column, str) else column)
    if not result:
        result = [{'source': 'id', 'label': 'ID'}]
    return result


def _cell(ds, context, column, labels=None):
    value = ds.value(context, column['source'])
    if column['source'].endswith('state') and labels:
        value = labels.get(value, value)
    if value is None or value == '' or value == []:
        return column.get('empty', '')
    return value if isinstance(value, (int, float, bool)) else _text(value)


def report_rows(conn, spec, view):
    ds, view = Dataset(conn, spec), _view(spec, view)
    columns = _columns(ds, view)
    contexts = [ctx for ctx in ds.contexts(view['entity']) if ds.matches(ctx, view.get('where', {}))]
    labels = view.get('workbook', {}).get('state_labels', {})
    return [c['label'] for c in columns], [[_cell(ds, ctx, col, labels) for col in columns] for ctx in contexts]


def _serialize(headers, rows, fmt, name, title=None, *, bom=False):
    if fmt == 'csv':
        stream = io.StringIO(newline='')
        writer = csv.writer(stream)
        writer.writerow([_safe(_text(value)) for value in headers])
        writer.writerows([[_safe(_text(value)) for value in row] for row in rows])
        return Download((('\ufeff' if bom else '') + stream.getvalue()).encode('utf-8'), 'text/csv', name + '.csv')
    if fmt == 'txt':
        if len(headers) != 1:
            raise ExportError('Text format requires one selected value column')
        return Download(('\n'.join(_text(row[0]) for row in rows) + ('\n' if rows else '')).encode('utf-8'), 'text/plain', name + '.txt')
    if fmt != 'xlsx':
        raise ExportError('Unsupported format')
    wb = Workbook()
    _table(wb.active, title or name, headers, rows)
    return _save(wb, name)


def _table(ws, title, headers, rows, *, labels=None, state_column=None, colors=None):
    _banner(ws, title, len(headers))
    _header(ws, headers)
    for index, values in enumerate(rows):
        _body_row(ws, index + 3, values, zebra=index % 2 == 1)
    if rows and state_column and labels:
        col = get_column_letter(state_column)
        target = col + '3:' + col + str(len(rows) + 2)
        # The closed spec validator enforces Excel's inline-list constraints.
        options = list(dict.fromkeys(labels.values()))
        if any(',' in item or '"' in item for item in options) or len(','.join(options)) > 253:
            raise ExportError('Dropdown labels must omit commas/quotes and fit 253 characters')
        _dropdown(ws, options, target)
        palette = {'positive': FILL_ALLOW, 'negative': FILL_DENY, 'waiting': FILL_PEND}
        for state, label in labels.items():
            fill = palette.get((colors or {}).get(state))
            if fill:
                ws.conditional_formatting.add(target, CellIsRule(operator='equal', formula=['"' + label + '"'], fill=fill))
    _widths(ws)
    ws.auto_filter.ref = 'A2:' + get_column_letter(len(headers)) + str(max(2, len(rows) + 2))


def _save(wb, name):
    stream = io.BytesIO()
    wb.save(stream)
    return Download(stream.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', name + '.xlsx')


def export_report(conn, spec, view, fmt):
    view = _view(spec, view)
    if fmt not in view.get('formats', ['csv', 'xlsx']):
        raise ExportError('Format is not declared for this report')
    if fmt == 'xlsx' and view.get('workbook'):
        return _workbook(Dataset(conn, spec), view)
    headers, rows = report_rows(conn, spec, view)
    return _serialize(headers, rows, fmt, view.get('name') or view['entity'], view.get('label'), bom=fmt == 'csv')


def declared_rows(conn, spec, name, *, text=False):
    ds = Dataset(conn, spec)
    matches = [item for item in ds.spec.get('exports', []) if item['name'] == name]
    if len(matches) != 1:
        raise ExportError('Unknown export')
    declaration = matches[0]
    root, *parts = declaration['from'].split('.')
    ds.entity(root)
    selector = '.'.join(parts)
    groups = {}
    for context in ds.contexts(root):
        if not ds.matches(context, declaration.get('where', {})):
            continue
        value = ds.value(context, selector) if selector else context['root']
        values = value if isinstance(value, list) else [value]
        for selected in values:
            if selected is None:
                continue
            item = dict(context, **{'$value': selected})
            key = _text(selected) if declaration.get('distinct') else str(len(groups))
            groups.setdefault(key, []).append(item)
    columns = declaration.get('columns')
    if text or not columns:
        if not selector:
            columns = [{'source': f['name'], 'label': f.get('label', f['name'])} for f in ds.entity(root)['fields']]
        else:
            columns = [{'source': '$value', 'label': parts[-1]}]
    rows = []
    for key in sorted(groups):
        contexts, row = groups[key], []
        for column in columns:
            values = [ds.value(ctx, column['source']) for ctx in contexts]
            aggregate = column.get('aggregate', 'value')
            if aggregate == 'join_distinct':
                if column.get('distinct_by'):
                    contributors = {}
                    for ctx, value in zip(contexts, values):
                        identity = _text(ds.value(ctx, column['distinct_by']))
                        if identity in contributors and contributors[identity] != _text(value):
                            raise ExportError('Contributor identity has conflicting display values')
                        contributors[identity] = _text(value)
                    value = '; '.join(value for _key, value in sorted(contributors.items(), key=lambda item: (item[1], item[0])))
                else:
                    value = '; '.join(sorted({_text(v) for v in values if v is not None}))
            elif aggregate == 'none_true':
                if any(type(v) is not bool for v in values):
                    raise ExportError('none_true requires boolean values')
                value = column.get('false_label', 'no') if any(values) else column.get('true_label', 'yes')
            elif aggregate == 'value':
                if any(v != values[0] for v in values):
                    raise ExportError('Grouped column requires an explicit aggregate')
                value = values[0]
            else:
                raise ExportError('Unknown aggregate')
            row.append(value)
        rows.append(row)
    return declaration, [column['label'] for column in columns], rows


def export_declared(conn, spec, name, fmt):
    declaration, headers, rows = declared_rows(conn, spec, name, text=fmt == 'txt')
    if fmt not in declaration['formats']:
        raise ExportError('Format is not declared for this export')
    return _serialize(headers, rows, fmt, name, declaration.get('label'))


def dashboard(conn, spec):
    ds, tiles = Dataset(conn, spec), []
    for tile in ds.spec.get('dashboard', []):
        entity = ds.entity(tile['count'])
        count = 0
        for row in ds.rows(entity['name']):
            context = {'root': row, 'entity': entity['name'], 'line': None}
            if ds.matches(context, tile.get('where', {}), record_state=True):
                count += 1
        tiles.append({'label': tile['label'], 'value': count})
    return tiles


def _workbook(ds, view):
    """Closed workbook layout: per-owner rows, touched-reference matrix, summary.

    All captions, column sources and state mappings are spec data. The sole
    formula is an engine COUNTIF over the generated matrix's owner cells.
    """
    cfg, columns = view['workbook'], _columns(ds, view)
    labels, matrix = cfg['state_labels'], cfg['matrix']
    contexts = [ctx for ctx in ds.contexts(view['entity']) if ds.matches(ctx, view.get('where', {}))]
    # Match the source's newest record per (owner, reference), not newest decision.
    contexts.sort(key=lambda ctx: (_text(ctx['root'].get('created_at')), str(ctx['root']['id'])), reverse=True)
    latest, owners, ticket_ids = {}, {}, {}
    for ctx in contexts:
        owner = str(ctx['root'].get('owner_id'))
        item = _text(ds.value(ctx, matrix['key']))
        person = ds.index['users'].get(owner)
        if not person or person.get('status') != 'approved':
            continue
        owners[owner] = person
        ticket_ids.setdefault(owner, set()).add(str(ctx['root']['id']))
        latest.setdefault((owner, item), ctx)
    owner_ids = sorted(owners, key=lambda key: (owners[key]['name'].casefold(), key))
    items = {}
    for (_owner, key), ctx in latest.items():
        items.setdefault(key, ctx)
    item_ids = sorted(items, key=lambda key: tuple(_text(ds.value(items[key], col['source'])) for col in matrix['columns']))
    wb, used = Workbook(), set()
    wb.remove(wb.active)
    # Reserve summary and matrix names before user names to prevent collisions.
    matrix_name = sheet_name_for(matrix['sheet'], used)
    summary_name = sheet_name_for(cfg['summary']['sheet'], used)
    state_source = cfg['state_source']
    status_column = next((i for i, col in enumerate(columns, 1) if col['source'] == state_source), None)
    for owner in owner_ids:
        person = owners[owner]
        hits = [latest[(owner, key)] for key in item_ids if (owner, key) in latest]
        ws = wb.create_sheet(sheet_name_for(person['name'], used))
        _table(ws, cfg['person_title'] + ' - ' + person['name'], [c['label'] for c in columns],
               [[_cell(ds, ctx, col, labels) for col in columns] for ctx in hits],
               labels=labels, state_column=status_column, colors=cfg.get('state_colors'))
    ws = wb.create_sheet(matrix_name)
    matrix_columns = matrix['columns']
    headers = [c['label'] for c in matrix_columns] + [owners[key]['name'] for key in owner_ids] + [matrix['total_label']]
    rows = []
    first = len(matrix_columns) + 1
    last = first + len(owner_ids) - 1
    matrix_labels = matrix['state_labels']
    count_value = matrix['count_value']
    if count_value not in matrix_labels.values() or any('"' in value for value in matrix_labels.values()):
        raise ExportError('Matrix count value must be a declared safe state label')
    for index, key in enumerate(item_ids, 3):
        values = [_cell(ds, items[key], col) for col in matrix_columns]
        for owner in owner_ids:
            ctx = latest.get((owner, key))
            state = ds.value(ctx, state_source) if ctx else None
            values.append(matrix_labels.get(state, '') if ctx else '')
        total = _Formula('=COUNTIF(' + get_column_letter(first) + str(index) + ':' + get_column_letter(last) + str(index) + ',"' + count_value + '")') if owner_ids else ''
        rows.append(values + [total])
    _table(ws, matrix['title'], headers, rows)
    if owner_ids and rows:
        target = get_column_letter(first) + '3:' + get_column_letter(last) + str(len(rows) + 2)
        choices = list(dict.fromkeys(matrix_labels.values()))
        if any(',' in item or '"' in item for item in choices) or len(','.join(choices)) > 253:
            raise ExportError('Invalid matrix dropdown labels')
        _dropdown(ws, choices, target)
        palette = {'positive': FILL_ALLOW, 'negative': FILL_DENY, 'waiting': FILL_PEND}
        for state, label in matrix_labels.items():
            if cfg.get('state_colors', {}).get(state) in palette:
                ws.conditional_formatting.add(target, CellIsRule(operator='equal', formula=['"' + label + '"'], fill=palette[cfg['state_colors'][state]]))
    summary = cfg['summary']
    rows = []
    for metric in summary['metrics']:
        kind = metric['kind']
        if kind == 'owners':
            value = len(owner_ids)
        elif kind == 'items':
            value = len(items)
        elif kind == 'records':
            value = sum(len(ids) for ids in ticket_ids.values())
        elif kind == 'groups':
            value = len({_text(ds.value(ctx, metric['source'])).split('>')[0].strip() for ctx in items.values()})
        elif kind == 'state':
            value = sum(ds.value(ctx, state_source) == metric['state'] for ctx in latest.values())
        else:
            raise ExportError('Unknown summary metric')
        rows.append([metric['label'], value])
    ws = wb.create_sheet(summary_name)
    ws.sheet_view.showGridLines = False
    _table(ws, summary['title'], summary['headers'], rows)
    if summary.get('per_person'):
        section = summary['per_person']
        start = len(rows) + 5
        headers = [section['name_label']] + [metric['label'] for metric in section['metrics']]
        ws.merge_cells(start_row=start - 1, start_column=1, end_row=start - 1, end_column=len(headers))
        cell = ws.cell(start - 1, 1, _safe(section['title']))
        cell.fill, cell.font = FILL_BANNER, FONT_BANNER
        cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        for n, header in enumerate(headers, 1):
            cell = ws.cell(start, n, _safe(header))
            cell.fill, cell.font = FILL_HEADER, FONT_HEADER
            cell.alignment = Alignment(horizontal='center' if n > 1 else 'left', vertical='center')
        for offset, owner in enumerate(owner_ids, 1):
            hits = [ctx for (key, _item), ctx in latest.items() if key == owner]
            values = [owners[owner]['name']]
            for metric in section['metrics']:
                kind = metric['kind']
                if kind == 'records':
                    value = len(ticket_ids[owner])
                elif kind == 'items':
                    value = len(hits)
                elif kind == 'groups':
                    value = len({_text(ds.value(ctx, metric['source'])).split('>')[0].strip() for ctx in hits})
                elif kind == 'state':
                    value = sum(ds.value(ctx, state_source) == metric['state'] for ctx in hits)
                elif kind == 'owners':
                    value = 1
                else:
                    raise ExportError('Unknown per-person metric')
                values.append(value)
            _body_row(ws, start + offset, values, zebra=offset % 2 == 0, center=tuple(range(2, len(headers) + 1)))
        _widths(ws)
    return _save(wb, view.get('name') or view['entity'])
